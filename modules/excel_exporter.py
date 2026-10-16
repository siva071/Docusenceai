"""In-memory Excel workbook generation for extracted PDF tables."""

from __future__ import annotations

from io import BytesIO
import json
import logging
import re
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill


class ExcelExportError(ValueError):
    """Raised when extracted tables cannot be represented as an Excel workbook."""


LOGGER = logging.getLogger(__name__)


_INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")
_MAX_SHEET_NAME_LENGTH = 31
_VALIDATION_SHEET_NAME = "Validation Report"
_SOURCE_EVIDENCE_SHEET_NAME = "Source Evidence"
_VALIDATION_COLUMNS = [
    "Table",
    "Source Page",
    "Row",
    "Related Row",
    "Column",
    "Severity",
    "Category",
    "Issue",
    "Suggested Action",
    "Evidence",
    "Actual Value",
    "Expected Value",
]


def _unique_sheet_name(candidate: object, used_names: set[str], fallback_index: int) -> str:
    """Return a valid, case-insensitively unique Excel worksheet name."""
    name = _INVALID_SHEET_CHARS.sub("_", str(candidate or "")).strip().strip("'")
    if not name:
        name = f"Table {fallback_index}"
    name = name[:_MAX_SHEET_NAME_LENGTH].strip().strip("'") or f"Table {fallback_index}"
    base_name = name
    suffix_number = 2
    while name.casefold() in used_names:
        suffix = f"_{suffix_number}"
        name = f"{base_name[:_MAX_SHEET_NAME_LENGTH - len(suffix)]}{suffix}"
        suffix_number += 1
    used_names.add(name.casefold())
    return name


def _table_to_frame(table: dict[str, Any], table_index: int) -> pd.DataFrame:
    """Build a rectangular dataframe while retaining every extracted row."""
    raw_rows = table.get("rows")
    if not isinstance(raw_rows, list):
        raise ExcelExportError(f"Table {table_index} has an invalid rows field.")

    records: list[tuple[list[Any], object]] = []
    for row_index, row in enumerate(raw_rows):
        if isinstance(row, dict):
            cells = row.get("cells", [])
            source_page = row.get("source_page", table.get("source_page"))
        elif isinstance(row, (list, tuple)):
            cells = list(row)
            source_page = table.get("source_page")
        else:
            raise ExcelExportError(f"Table {table_index}, row {row_index + 1} is malformed.")
        if not isinstance(cells, (list, tuple)):
            raise ExcelExportError(f"Table {table_index}, row {row_index + 1} has invalid cells.")
        records.append((list(cells), source_page))

    headers_value = table.get("headers")
    inferred_headers = (
        isinstance(headers_value, (list, tuple))
        and bool(headers_value)
        and bool(table.get("headers_inferred", True))
    )
    headers = list(headers_value) if inferred_headers else []
    if records and records[0][0] and raw_rows and isinstance(raw_rows[0], dict):
        if raw_rows[0].get("is_inferred_header") and records[0][0] == headers:
            records.pop(0)

    column_count = max(
        [len(headers), int(table.get("column_count", 0) or 0)]
        + [len(cells) for cells, _ in records]
    )
    if column_count < 1:
        raise ExcelExportError(
            f"Table {table_index} is empty and has no headers or columns to export."
        )

    headers.extend(f"Column {index + 1}" for index in range(len(headers), column_count))
    headers = headers[:column_count]

    source_page_header = "Source Page"
    while source_page_header in headers:
        source_page_header = f"{source_page_header} (metadata)"
    output_headers = headers + [source_page_header]

    output_rows: list[list[Any]] = []
    for cells, source_page in records:
        # Pad short rows but never drop non-empty or partially populated data.
        padded_cells = cells[:column_count] + [None] * max(0, column_count - len(cells))
        output_rows.append(padded_cells + [source_page])

    return pd.DataFrame(output_rows, columns=output_headers)


def _source_evidence_rows(tables: list[dict[str, Any]]) -> list[list[Any]]:
    """Flatten per-cell provenance, explicitly retaining unavailable evidence."""
    evidence_rows: list[list[Any]] = []
    for table_index, table in enumerate(tables, start=1):
        table_name = str(table.get("table_id") or f"Table {table_index}")
        headers_value = table.get("headers")
        headers = list(headers_value) if isinstance(headers_value, (list, tuple)) else []
        for row_index, row in enumerate(table.get("rows", []), start=1):
            if isinstance(row, dict):
                cells = row.get("cells", [])
                source_page = row.get("source_page", table.get("source_page"))
                source_row = row.get("source_row_index", row_index)
                default_method = row.get("extraction_method", table.get("extraction_method"))
                cell_evidence = row.get("cell_evidence", [])
            elif isinstance(row, (list, tuple)):
                cells = list(row)
                source_page = table.get("source_page")
                source_row = row_index
                default_method = table.get("extraction_method")
                cell_evidence = []
            else:
                continue
            if not isinstance(cells, (list, tuple)):
                continue
            for column_index, value in enumerate(cells, start=1):
                evidence = (
                    cell_evidence[column_index - 1]
                    if isinstance(cell_evidence, (list, tuple))
                    and column_index <= len(cell_evidence)
                    and isinstance(cell_evidence[column_index - 1], dict)
                    else {}
                )
                evidence_status = str(evidence.get("evidence_status") or "unavailable")
                bbox = evidence.get("bbox")
                source_text = evidence.get("source_text")
                word_boxes = evidence.get("word_boxes") or []
                method = evidence.get("extraction_method") or default_method or "unknown"
                limitation = evidence.get("confidence_or_limitations") or (
                    "Exact cell-level source evidence is unavailable; source page and extracted position are the only known context."
                )
                evidence_rows.append(
                    [
                        table_name,
                        row_index,
                        source_row,
                        f"{headers[column_index - 1]} (column {column_index})"
                        if column_index <= len(headers)
                        else f"Column {column_index}",
                        evidence.get("source_column"),
                        value,
                        evidence.get("source_page", source_page),
                        method,
                        evidence.get("source_text_method")
                        or (
                            "ocr" if method == "ocr" and source_text
                            else "digital_text" if method == "pdfplumber" and source_text
                            else "unavailable"
                        ),
                        evidence.get("ocr_confidence"),
                        evidence_status,
                        json.dumps(bbox, ensure_ascii=False) if bbox is not None else None,
                        source_text,
                        json.dumps(word_boxes, ensure_ascii=False) if word_boxes else None,
                        limitation,
                    ]
                )
    return evidence_rows


def export_tables_to_excel(
    tables: list[dict[str, Any]],
    validation_report: dict[str, Any] | None = None,
) -> bytes:
    """Create an `.xlsx` workbook with table, validation, and evidence worksheets.

    Table dictionaries use the extraction engine's schema: a ``table_id``,
    optional inferred ``headers``, and row records containing ``cells`` and
    ``source_page``. Inferred header rows are written as worksheet headers;
    all remaining rows are kept in source order, including partial rows.
    The workbook is built entirely in memory.

    Raises:
        ExcelExportError: If no tables were detected or workbook creation fails.
    """
    if not isinstance(tables, list) or not tables:
        raise ExcelExportError("No tables were detected. There is no workbook to export.")

    if validation_report is None:
        validation_report = {
            "findings": [],
            "summary": "Validation was not run for this workbook.",
        }
    findings = validation_report.get("findings", [])
    if not isinstance(findings, list):
        raise ExcelExportError("Validation report findings must be a list.")

    output = BytesIO()
    try:
        # Reserve the report sheet name before assigning extracted table names.
        used_names: set[str] = {
            _VALIDATION_SHEET_NAME.casefold(),
            _SOURCE_EVIDENCE_SHEET_NAME.casefold(),
        }
        prepared_sheets: list[tuple[str, pd.DataFrame]] = []
        for table_index, table in enumerate(tables, start=1):
            if not isinstance(table, dict):
                raise ExcelExportError(f"Table {table_index} has an invalid format.")
            frame = _table_to_frame(table, table_index)
            sheet_name = _unique_sheet_name(
                table.get("table_id") or f"Table {table_index}",
                used_names,
                table_index,
            )
            prepared_sheets.append((sheet_name, frame))

        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            for sheet_name, frame in prepared_sheets:
                frame.to_excel(writer, sheet_name=sheet_name, index=False)

        # Use openpyxl for workbook-level polish and force all PDF-derived
        # strings to remain literal cell text rather than spreadsheet formulas.
        workbook = load_workbook(BytesIO(output.getvalue()))
        report_sheet = workbook.create_sheet(_VALIDATION_SHEET_NAME)
        report_sheet.append(_VALIDATION_COLUMNS)
        if findings:
            for finding in findings:
                if not isinstance(finding, dict):
                    raise ExcelExportError("Validation report contains a malformed finding.")
                report_sheet.append(
                    [
                        finding.get("table_name"),
                        finding.get("source_page"),
                        finding.get("row"),
                        finding.get("related_row"),
                        finding.get("column"),
                        finding.get("severity"),
                        finding.get("category"),
                        finding.get("issue"),
                        finding.get("suggested_action"),
                        finding.get("evidence"),
                        finding.get("actual_value"),
                        finding.get("expected_value"),
                    ]
                )
        else:
            report_sheet.append(
                [
                    str(validation_report.get("summary") or "No validation issues were detected by the implemented checks."),
                    *([None] * (len(_VALIDATION_COLUMNS) - 1)),
                ]
            )
        evidence_sheet = workbook.create_sheet(_SOURCE_EVIDENCE_SHEET_NAME)
        evidence_sheet.append(
            [
                "Table",
                "Table Row",
                "Source Row",
                "Table Column",
                "Source Column",
                "Value",
                "Page Number",
                "Extraction Method",
                "Source Text Method",
                "OCR Confidence",
                "Evidence Status",
                "Bounding Box (PDF points)",
                "Source Text",
                "Word Boxes",
                "Confidence / Limitations",
            ]
        )
        evidence_rows = _source_evidence_rows(tables)
        if evidence_rows:
            for evidence_row in evidence_rows:
                evidence_sheet.append(evidence_row)
        else:
            evidence_sheet.append(
                ["No extracted cell values were available for source evidence."]
                + [None] * 13
            )
        for worksheet in workbook.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for cell in worksheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(fill_type="solid", fgColor="147D78")
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f" and isinstance(cell.value, str):
                        cell.data_type = "s"
            for column_cells in worksheet.columns:
                column_letter = column_cells[0].column_letter
                width = max(
                    (len(str(cell.value)) for cell in column_cells if cell.value is not None),
                    default=10,
                )
                worksheet.column_dimensions[column_letter].width = min(max(width + 2, 12), 55)

        finalized = BytesIO()
        workbook.save(finalized)
        return finalized.getvalue()
    except ExcelExportError:
        raise
    except Exception as exc:
        LOGGER.exception("Excel workbook generation failed")
        raise ExcelExportError(
            "Excel workbook generation failed. Review the extracted tables and try again."
        ) from exc
