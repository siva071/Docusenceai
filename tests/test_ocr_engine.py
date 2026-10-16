"""Mocked local OCR tests using synthetic image-only and mixed PDFs."""

from __future__ import annotations

from io import BytesIO
import unittest
from unittest.mock import patch
import builtins

from PIL import Image, ImageDraw

try:
    import pymupdf as fitz
except ImportError:
    import fitz

from modules.ocr_engine import check_ocr_runtime, is_likely_scanned_page
from modules.pdf_extractor import extract_pdf_tables, extract_pdf_text, retry_failed_ocr_pages
from modules.table_stitcher import stitch_tables
from modules.excel_exporter import export_tables_to_excel
from openpyxl import load_workbook
import pytesseract

_ORIGINAL_IMPORT = builtins.__import__


def _image_bytes() -> bytes:
    """Return a synthetic scanned page image with visible text."""
    image = Image.new("RGB", (1000, 1400), "white")
    draw = ImageDraw.Draw(image)
    draw.text((100, 100), "Scanned source page", fill="black")
    output = BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()


def _pdf_with_scanned_pages(page_count: int = 1, scanned_pages: set[int] | None = None) -> bytes:
    """Build a synthetic PDF with embedded raster pages and selectable pages."""
    scanned = scanned_pages or {1}
    document = fitz.open()
    for page_number in range(1, page_count + 1):
        page = document.new_page(width=612, height=792)
        if page_number in scanned:
            page.insert_image(page.rect, stream=_image_bytes())
        else:
            page.insert_text((72, 72), f"Digital text on page {page_number} with enough selectable characters")
    pdf_bytes = document.tobytes()
    document.close()
    return pdf_bytes


def _ocr_data(rows: list[list[str]], x_positions: list[int] | None = None) -> dict[str, list[object]]:
    """Return a pytesseract-shaped TSV dictionary for controlled OCR output."""
    if x_positions is None:
        x_positions = [100 + index * 180 for index in range(len(rows[0]))]
    data: dict[str, list[object]] = {
        "text": [],
        "conf": [],
        "left": [],
        "top": [],
        "width": [],
        "height": [],
        "block_num": [],
        "par_num": [],
        "line_num": [],
    }
    for line_index, row in enumerate(rows, start=1):
        for column_index, value in enumerate(row):
            if value == "":
                continue
            data["text"].append(value)
            data["conf"].append("92.0")
            data["left"].append(x_positions[column_index])
            data["top"].append(300 + (line_index - 1) * 120)
            data["width"].append(max(24, len(value) * 12))
            data["height"].append(30)
            data["block_num"].append(1)
            data["par_num"].append(1)
            data["line_num"].append(line_index)
    return data


class OcrDetectionTests(unittest.TestCase):
    def test_detects_image_heavy_low_text_page_but_not_digital_page(self) -> None:
        self.assertTrue(is_likely_scanned_page("", 0.9))
        self.assertTrue(is_likely_scanned_page("a few words", 0.5))
        self.assertFalse(is_likely_scanned_page("This page has enough selectable text", 0.9))
        self.assertFalse(is_likely_scanned_page("", 0.01))

    @patch("pytesseract.image_to_data")
    def test_ocr_extracts_values_and_page_metadata(self, mocked_ocr) -> None:
        mocked_ocr.return_value = _ocr_data([["Invoice", "INV-204"], ["Total", "-1,234.50"]])
        pdf_bytes = _pdf_with_scanned_pages()

        result = extract_pdf_text(pdf_bytes)

        page = result["pages"][0]
        self.assertEqual(page["page_number"], 1)
        self.assertEqual(page["status"], "ocr_text")
        self.assertEqual(page["extraction_method"], "ocr")
        self.assertIn("INV-204", page["text"])
        self.assertIn("-1,234.50", page["text"])
        self.assertTrue(page["ocr_attempted"])
        self.assertEqual(page["ocr_metadata"]["engine"], "Tesseract")
        self.assertTrue(all(word["page_number"] == 1 for word in page["ocr_words"]))

    @patch("pytesseract.image_to_data")
    def test_mixed_pdf_ocr_only_runs_for_scanned_page(self, mocked_ocr) -> None:
        mocked_ocr.return_value = _ocr_data([["OCR", "only page two"]])
        pdf_bytes = _pdf_with_scanned_pages(page_count=2, scanned_pages={2})

        result = extract_pdf_text(pdf_bytes)

        self.assertEqual(mocked_ocr.call_count, 1)
        self.assertEqual(result["pages"][0]["status"], "digital_text")
        self.assertFalse(result["pages"][0]["ocr_attempted"])
        self.assertEqual(result["pages"][1]["page_number"], 2)
        self.assertEqual(result["pages"][1]["status"], "ocr_text")
        self.assertIn("OCR only page two", result["pages"][1]["text"])

    @patch("pytesseract.image_to_data")
    def test_ocr_failure_is_page_specific_and_later_pages_continue(self, mocked_ocr) -> None:
        mocked_ocr.side_effect = [RuntimeError("simulated OCR failure"), _ocr_data([["Last", "page"]])]
        pdf_bytes = _pdf_with_scanned_pages(page_count=2, scanned_pages={1, 2})

        result = extract_pdf_text(pdf_bytes)

        self.assertEqual(result["pages"][0]["status"], "ocr_failed")
        self.assertIn("developer diagnostics", result["pages"][0]["error"])
        self.assertEqual(result["pages"][1]["status"], "ocr_text")
        self.assertIn("Last page", result["pages"][1]["text"])
        self.assertTrue(any("Page 1" in warning for warning in result["warnings"]))

    @patch("pytesseract.image_to_data")
    def test_missing_tesseract_reports_setup_help(self, mocked_ocr) -> None:
        mocked_ocr.side_effect = pytesseract.TesseractNotFoundError()
        result = extract_pdf_text(_pdf_with_scanned_pages())

        self.assertEqual(result["pages"][0]["status"], "ocr_failed")
        self.assertIn("TESSERACT_CMD", result["pages"][0]["error"])

    @patch("builtins.__import__")
    def test_missing_python_ocr_dependency_is_actionable(self, mocked_import) -> None:
        def import_without_pytesseract(name, *args, **kwargs):
            if name == "pytesseract":
                raise ImportError("simulated missing pytesseract")
            return _ORIGINAL_IMPORT(name, *args, **kwargs)

        mocked_import.side_effect = import_without_pytesseract
        result = check_ocr_runtime()

        self.assertFalse(result["available"])
        self.assertIn("Install project requirements", result["message"])

    @patch("builtins.__import__")
    def test_missing_pillow_dependency_is_actionable(self, mocked_import) -> None:
        def import_without_pillow(name, *args, **kwargs):
            if name == "PIL.Image":
                raise ImportError("simulated missing Pillow")
            return _ORIGINAL_IMPORT(name, *args, **kwargs)

        mocked_import.side_effect = import_without_pillow
        result = check_ocr_runtime()

        self.assertFalse(result["available"])
        self.assertIn("image dependency is missing", result["message"])

    @patch("pytesseract.image_to_data")
    def test_retry_ocr_success_only_reprocesses_failed_page_and_preserves_others(self, mocked_ocr) -> None:
        mocked_ocr.side_effect = [
            _ocr_data([["Success", "page-one"]]),
            RuntimeError("temporary page-two failure"),
            RuntimeError("page-three failure remains"),
            _ocr_data([["Recovered", "page-two"]]),
        ]
        pdf_bytes = _pdf_with_scanned_pages(page_count=3, scanned_pages={1, 2, 3})
        initial = extract_pdf_text(pdf_bytes)
        original_success = dict(initial["pages"][0])
        original_third_page = dict(initial["pages"][2])
        self.assertEqual([page["ocr_status"] for page in initial["pages"]], ["ocr_text", "ocr_failed", "ocr_failed"])

        retried = retry_failed_ocr_pages(pdf_bytes, initial["pages"], page_numbers=[2])

        self.assertEqual(retried["retried_pages"], [2])
        self.assertEqual(mocked_ocr.call_count, 4)
        self.assertEqual(retried["pages"][0], original_success)
        self.assertEqual(retried["pages"][2], original_third_page)
        self.assertEqual(retried["pages"][1]["ocr_status"], "ocr_text")
        self.assertIn("Recovered page-two", retried["pages"][1]["text"])

    @patch("pytesseract.image_to_data")
    def test_retry_failure_remains_failed_and_successful_pages_are_unchanged(self, mocked_ocr) -> None:
        mocked_ocr.side_effect = [
            _ocr_data([["Already", "good"]]),
            RuntimeError("initial OCR failure"),
            RuntimeError("retry OCR failure"),
        ]
        pdf_bytes = _pdf_with_scanned_pages(page_count=2, scanned_pages={1, 2})
        initial = extract_pdf_text(pdf_bytes)
        page_one_before = dict(initial["pages"][0])

        retried = retry_failed_ocr_pages(pdf_bytes, initial["pages"], page_numbers=[2])

        self.assertEqual(retried["pages"][0], page_one_before)
        self.assertEqual(retried["pages"][1]["ocr_status"], "ocr_failed")
        self.assertIn("developer diagnostics", retried["pages"][1]["error"])
        self.assertEqual(mocked_ocr.call_count, 3)

    @patch("pytesseract.image_to_data")
    def test_ocr_layout_table_supports_ten_columns_and_source_page(self, mocked_ocr) -> None:
        headers = [f"Field{index}" for index in range(1, 11)]
        values = ["ID-01", "-12.50", "1,234.56", "2026-10-16", "A5", "06", "7.00", "8", "9", "10"]
        partial_values = ["ID-02", "", "0.25", "2026-10-17", "B7", "11", "12", "13", "14", "15"]
        mocked_ocr.return_value = _ocr_data([headers, values, partial_values])
        pdf_bytes = _pdf_with_scanned_pages()

        text_result = extract_pdf_text(pdf_bytes)
        tables_result = extract_pdf_tables(pdf_bytes, ocr_pages=text_result["pages"])

        self.assertEqual(len(tables_result["tables"]), 1)
        table = tables_result["tables"][0]
        self.assertEqual(table["extraction_method"], "ocr_layout")
        self.assertEqual(table["column_count"], 10)
        self.assertEqual(table["headers"], headers)
        self.assertEqual(table["rows"][1]["cells"], values)
        self.assertEqual(table["rows"][2]["cells"][0], "ID-02")
        self.assertEqual(table["rows"][2]["cells"][1], "")
        self.assertEqual([row["source_page"] for row in table["rows"]], [1, 1, 1])
        cell_evidence = table["rows"][1]["cell_evidence"][1]
        self.assertEqual(cell_evidence["source_page"], 1)
        self.assertEqual(cell_evidence["extraction_method"], "ocr")
        self.assertEqual(cell_evidence["source_text"], "-12.50")
        self.assertTrue(cell_evidence["word_boxes"])
        self.assertEqual(cell_evidence["evidence_status"], "source_text_mapped")
        blank_evidence = table["rows"][2]["cell_evidence"][1]
        self.assertIsNone(blank_evidence["source_text"])
        self.assertEqual(blank_evidence["evidence_status"], "unavailable")
        self.assertEqual(table["rows"][1]["cells"][1], "-12.50")
        self.assertEqual(table["rows"][1]["cells"][2], "1,234.56")

        stitched = stitch_tables(tables_result["tables"])
        workbook = load_workbook(BytesIO(export_tables_to_excel(stitched.tables)), data_only=True)
        worksheet = workbook[workbook.sheetnames[0]]
        self.assertEqual(worksheet.cell(row=2, column=1).value, "ID-01")
        self.assertEqual(worksheet.cell(row=2, column=2).value, "-12.50")
        self.assertEqual(worksheet.cell(row=2, column=3).value, "1,234.56")
        self.assertEqual(worksheet.cell(row=2, column=11).value, 1)
        evidence_sheet = workbook["Source Evidence"]
        evidence_row = next(
            row
            for row in evidence_sheet.iter_rows(min_row=2, values_only=True)
            if row[5] == "-12.50"
        )
        self.assertEqual(evidence_row[6], 1)
        self.assertEqual(evidence_row[7], "ocr")
        self.assertEqual(evidence_row[8], "ocr")
        self.assertEqual(evidence_row[9], 92.0)
        self.assertEqual(evidence_row[10], "source_text_mapped")
        self.assertEqual(evidence_row[12], "-12.50")
        self.assertIn('"bbox"', evidence_row[13])


if __name__ == "__main__":
    unittest.main()
