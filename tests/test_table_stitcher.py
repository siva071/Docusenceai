"""Synthetic multi-page PDF tests for cautious table stitching."""

from __future__ import annotations

from io import BytesIO
import unittest

from openpyxl import load_workbook

try:
    import pymupdf as fitz
except ImportError:
    import fitz

from modules.excel_exporter import export_tables_to_excel
from modules.pdf_extractor import extract_pdf_tables
from modules.table_stitcher import stitch_tables

PAGE_WIDTH = 612
PAGE_HEIGHT = 792


def _draw_ruled_table(
    page: fitz.Page,
    rows: list[list[str]],
    x_edges: list[float],
    top: float,
    row_height: float = 34,
) -> None:
    """Draw a ruled table with selectable cell text into a PDF page."""
    bottom = top + len(rows) * row_height
    for x in x_edges:
        page.draw_line((x, top), (x, bottom), width=0.8)
    for row_index in range(len(rows) + 1):
        y = top + row_index * row_height
        page.draw_line((x_edges[0], y), (x_edges[-1], y), width=0.8)
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            if value:
                page.insert_text(
                    (x_edges[column_index] + 5, top + row_index * row_height + 23),
                    value,
                    fontsize=9,
                )


def make_multipage_pdf(page_tables: list[list[dict[str, object]]]) -> bytes:
    """Create a PDF from per-page table fixtures; empty lists produce blank pages."""
    document = fitz.open()
    for page_tables_for_page in page_tables:
        page = document.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)
        for table in page_tables_for_page:
            _draw_ruled_table(
                page,
                table["rows"],
                table["x_edges"],
                table["top"],
                table.get("row_height", 34),
            )
    pdf_bytes = document.tobytes()
    document.close()
    return pdf_bytes


def table(rows: list[list[str]], top: float, x_edges: list[float] | None = None) -> dict[str, object]:
    """Build a positioned table fixture."""
    return {"rows": rows, "top": top, "x_edges": x_edges or [72, 306, 540]}


class TableStitcherTests(unittest.TestCase):
    def test_stitches_repeat_headers_with_formatting_differences_and_exports_rows(self) -> None:
        pdf_bytes = make_multipage_pdf(
            [
                [
                    table(
                        [["Account ID", "Amount"], ["A-1", "10"], ["A-2", ""]],
                        top=620,
                    )
                ],
                [
                    table(
                        [[" account-id ", " AMOUNT "], ["A-3", "30"], ["A-4", "40"]],
                        top=48,
                    )
                ],
            ]
        )

        extracted = extract_pdf_tables(pdf_bytes)
        result = stitch_tables(extracted["tables"])

        self.assertEqual(len(extracted["tables"]), 2)
        self.assertEqual(len(result.tables), 1)
        self.assertEqual(result.stitched_count, 1)
        stitched = result.tables[0]
        self.assertEqual(stitched["source_pages"], [1, 2])
        self.assertEqual(
            [row["cells"] for row in stitched["rows"]],
            [
                ["Account ID", "Amount"],
                ["A-1", "10"],
                ["A-2", ""],
                ["A-3", "30"],
                ["A-4", "40"],
            ],
        )
        self.assertEqual(
            [row["source_page"] for row in stitched["rows"]],
            [1, 1, 1, 2, 2],
        )
        self.assertEqual(stitched["rows"][1]["cell_evidence"][0]["source_page"], 1)
        self.assertEqual(stitched["rows"][3]["cell_evidence"][0]["source_page"], 2)
        self.assertEqual(stitched["rows"][3]["source_row_index"], 2)

        workbook = load_workbook(BytesIO(export_tables_to_excel(result.tables)), data_only=True)
        worksheet = workbook[workbook.sheetnames[0]]
        self.assertEqual(worksheet.max_row, 5)
        self.assertEqual([cell.value for cell in worksheet[1]], ["Account ID", "Amount", "Source Page"])
        self.assertEqual([worksheet.cell(row=2, column=1).value, worksheet.cell(row=3, column=1).value], ["A-1", "A-2"])
        self.assertEqual([worksheet.cell(row=row, column=3).value for row in range(2, 6)], [1, 1, 2, 2])
        self.assertNotIn(" account-id ", [worksheet.cell(row=row, column=1).value for row in range(2, 6)])

    def test_stitches_continuation_without_repeated_header(self) -> None:
        pdf_bytes = make_multipage_pdf(
            [
                [table([["Code", "Description"], ["01", "First"]], top=625)],
                [table([["02", "Second"], ["03", "Third"]], top=50)],
            ]
        )

        result = stitch_tables(extract_pdf_tables(pdf_bytes)["tables"])

        self.assertEqual(len(result.tables), 1)
        self.assertEqual(result.stitched_count, 1)
        self.assertEqual(
            [row["cells"] for row in result.tables[0]["rows"]],
            [["Code", "Description"], ["01", "First"], ["02", "Second"], ["03", "Third"]],
        )
        self.assertEqual([row["source_page"] for row in result.tables[0]["rows"]], [1, 1, 2, 2])

    def test_unrelated_tables_with_same_column_count_stay_separate(self) -> None:
        pdf_bytes = make_multipage_pdf(
            [
                [table([["Account", "Amount"], ["A-1", "10"]], top=625)],
                [table([["Location", "Owner"], ["North", "Sam"]], top=50)],
            ]
        )

        result = stitch_tables(extract_pdf_tables(pdf_bytes)["tables"])

        self.assertEqual(len(result.tables), 2)
        self.assertEqual(result.stitched_count, 0)
        self.assertTrue(any("different headers" in warning for warning in result.warnings))
        self.assertEqual(result.tables[0]["rows"][-1]["cells"], ["A-1", "10"])
        self.assertEqual(result.tables[1]["rows"][-1]["cells"], ["North", "Sam"])

    def test_different_column_structures_stay_separate_with_warning(self) -> None:
        pdf_bytes = make_multipage_pdf(
            [
                [table([["Account", "Amount"], ["A-1", "10"]], top=625)],
                [table([["Account", "Amount", "Status"], ["A-2", "20", "Open"]], top=50,
                       x_edges=[72, 228, 384, 540])],
            ]
        )

        result = stitch_tables(extract_pdf_tables(pdf_bytes)["tables"])

        self.assertEqual(len(result.tables), 2)
        self.assertTrue(any("different column counts" in warning for warning in result.warnings))

    def test_tables_separated_by_a_blank_page_do_not_stitch(self) -> None:
        pdf_bytes = make_multipage_pdf(
            [
                [table([["Code", "Value"], ["A", "1"]], top=625)],
                [],
                [table([["Code", "Value"], ["B", "2"]], top=50)],
            ]
        )

        result = stitch_tables(extract_pdf_tables(pdf_bytes)["tables"])

        self.assertEqual(len(result.tables), 2)
        self.assertEqual(result.stitched_count, 0)
        self.assertEqual([entry["source_page"] for entry in result.tables], [1, 3])
        self.assertEqual([entry["rows"][-1]["cells"] for entry in result.tables], [["A", "1"], ["B", "2"]])

    def test_three_page_continuation_retains_all_source_pages(self) -> None:
        middle_rows = [[f"Middle-{index}", str(index)] for index in range(18)]
        last_rows = [["C", "3"], ["D", "4"]]
        pdf_bytes = make_multipage_pdf(
            [
                [table([["Code", "Value"], ["A", "1"]], top=625)],
                [table(middle_rows, top=50)],
                [table(last_rows, top=50)],
            ]
        )

        result = stitch_tables(extract_pdf_tables(pdf_bytes)["tables"])

        self.assertEqual(len(result.tables), 1)
        self.assertEqual(result.stitched_count, 2)
        self.assertEqual(result.tables[0]["source_pages"], [1, 2, 3])
        self.assertEqual(result.tables[0]["rows"][1]["cells"], ["A", "1"])
        self.assertEqual(result.tables[0]["rows"][-2]["cells"], ["C", "3"])
        self.assertEqual([row["source_page"] for row in result.tables[0]["rows"]].count(1), 2)
        self.assertEqual([row["source_page"] for row in result.tables[0]["rows"]].count(2), 18)
        self.assertEqual([row["source_page"] for row in result.tables[0]["rows"]].count(3), 2)

    def test_repeated_header_and_data_fragment_is_removed_and_continues(self) -> None:
        page_two_rows = [["Reference", "Value"]] + [
            [f"R-{index}", str(index)] for index in range(1, 19)
        ]
        pdf_bytes = make_multipage_pdf(
            [
                [table([["Reference", "Value"], ["R-1", "10"]], top=625)],
                [table(page_two_rows, top=50)],
                [table([["R-19", "19"], ["R-20", "20"]], top=50)],
            ]
        )

        result = stitch_tables(extract_pdf_tables(pdf_bytes)["tables"])

        self.assertEqual(len(result.tables), 1)
        self.assertEqual(result.stitched_count, 2)
        all_rows = [row["cells"] for row in result.tables[0]["rows"]]
        self.assertEqual(all_rows[0], ["Reference", "Value"])
        self.assertEqual(all_rows[1], ["R-1", "10"])
        self.assertNotIn(["Reference", "Value"], all_rows[2:])
        self.assertEqual(all_rows[-2:], [["R-19", "19"], ["R-20", "20"]])
        source_pages = [row["source_page"] for row in result.tables[0]["rows"]]
        self.assertEqual(source_pages.count(1), 2)
        self.assertEqual(source_pages.count(2), 18)
        self.assertEqual(source_pages.count(3), 2)


if __name__ == "__main__":
    unittest.main()
