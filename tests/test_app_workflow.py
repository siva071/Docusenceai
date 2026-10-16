"""Streamlit integration test for upload, extraction, and Excel download UI."""

from __future__ import annotations

from pathlib import Path
import unittest
from io import BytesIO
from unittest.mock import patch

from PIL import Image
from openpyxl import load_workbook

try:
    import pymupdf as fitz
except ImportError:
    import fitz

from streamlit.testing.v1 import AppTest

PROJECT_ROOT = Path(__file__).resolve().parents[1]


def make_table_pdf() -> bytes:
    """Create a one-page ruled table with known header and row values."""
    document = fitz.open()
    page = document.new_page(width=612, height=792)
    rows = [["Account", "Balance"], ["Cash", "100.00"]]
    x0, y0, cell_width, cell_height = 72, 100, 180, 40
    for column in range(3):
        x = x0 + column * cell_width
        page.draw_line((x, y0), (x, y0 + len(rows) * cell_height), width=0.8)
    for row_index in range(len(rows) + 1):
        y = y0 + row_index * cell_height
        page.draw_line((x0, y), (x0 + 2 * cell_width, y), width=0.8)
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            page.insert_text(
                (x0 + column_index * cell_width + 5, y0 + row_index * cell_height + 25),
                value,
            )
    pdf_bytes = document.tobytes()
    document.close()
    return pdf_bytes


def make_text_only_pdf() -> bytes:
    """Create a digital PDF page that contains no table."""
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Readable digital text without a table")
    pdf_bytes = document.tobytes()
    document.close()
    return pdf_bytes


def make_scanned_pdf(page_count: int = 1) -> bytes:
    """Create image-only PDF pages for mocked OCR UI integration."""
    image = Image.new("RGB", (300, 400), "white")
    image_bytes = BytesIO()
    image.save(image_bytes, format="PNG")
    document = fitz.open()
    for _ in range(page_count):
        page = document.new_page(width=612, height=792)
        page.insert_image(page.rect, stream=image_bytes.getvalue())
    pdf_bytes = document.tobytes()
    document.close()
    return pdf_bytes


def make_validation_pdf() -> bytes:
    """Create a ruled table with a numeric-text anomaly and a suspicious blank."""
    document = fitz.open()
    page = document.new_page(width=612, height=792)
    rows = [["Account", "Amount", "Status"], ["A-1", "10.00", "Paid"], ["A-2", "bad", ""]]
    x_edges = [72, 220, 368, 516]
    y0, row_height = 120, 42
    for x in x_edges:
        page.draw_line((x, y0), (x, y0 + len(rows) * row_height), width=0.8)
    for row_index in range(len(rows) + 1):
        y = y0 + row_index * row_height
        page.draw_line((x_edges[0], y), (x_edges[-1], y), width=0.8)
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            if value:
                page.insert_text(
                    (x_edges[column_index] + 5, y0 + row_index * row_height + 27),
                    value,
                )
    pdf_bytes = document.tobytes()
    document.close()
    return pdf_bytes


class AppWorkflowTests(unittest.TestCase):
    def test_pdf_upload_extraction_and_download_button(self) -> None:
        app = AppTest.from_file(str(PROJECT_ROOT / "app.py")).run()
        app.file_uploader[0].upload(
            "workflow-test.pdf", make_table_pdf(), "application/pdf"
        ).run()
        self.assertEqual(len(app.exception), 0)
        self.assertTrue(any(button.label == "Start Processing" for button in app.button))

        next(button for button in app.button if button.label == "Start Processing").click().run()

        self.assertEqual(len(app.exception), 0)
        table_metric = next(metric for metric in app.metric if metric.label == "Tables after stitching")
        self.assertEqual(table_metric.value, "1")
        self.assertEqual(len(app.dataframe), 2)  # table preview and per-page recovery status
        table_preview = next(
            frame for frame in app.dataframe if "Account" in frame.value.columns
        )
        self.assertEqual(
            list(table_preview.value.columns),
            ["Source page", "Source row", "Account", "Balance"],
        )
        self.assertTrue(any(metric.label == "Rows extracted" and metric.value == "1" for metric in app.metric))
        download_buttons = app.get("download_button")
        self.assertEqual(len(download_buttons), 1)
        self.assertEqual(download_buttons[0].label, "Download Excel")
        extraction_result = app.session_state["extraction_result"]
        self.assertIsNotNone(extraction_result["validation"])
        workbook = load_workbook(BytesIO(extraction_result["excel_bytes"]), data_only=True)
        self.assertIn("Validation Report", workbook.sheetnames)
        self.assertIn("Source Evidence", workbook.sheetnames)
        self.assertIn("No validation issues", workbook["Validation Report"]["A2"].value)
        self.assertEqual(workbook["Source Evidence"]["A1"].value, "Table")
        self.assertGreater(workbook["Source Evidence"].max_row, 1)
        self.assertEqual(workbook["page-1-table-1"].max_row, 2)
        self.assertEqual(workbook["page-1-table-1"].max_column, 3)
        self.assertEqual(workbook["page-1-table-1"]["B2"].value, "100.00")

    def test_proof_mode_selects_cell_and_displays_evidence_without_errors(self) -> None:
        app = AppTest.from_file(str(PROJECT_ROOT / "app.py")).run()
        app.file_uploader[0].upload("proof.pdf", make_table_pdf(), "application/pdf").run()
        next(button for button in app.button if button.label == "Start Processing").click().run()

        self.assertEqual(len(app.exception), 0)
        proof_selectboxes = [
            selectbox
            for selectbox in app.selectbox
            if selectbox.label in {"Table", "Extracted row", "Cell"}
        ]
        self.assertEqual({selectbox.label for selectbox in proof_selectboxes}, {"Table", "Extracted row", "Cell"})
        self.assertTrue(any("Extracted value" in markdown.value for markdown in app.markdown))

    def test_no_table_upload_shows_message_and_no_download(self) -> None:
        app = AppTest.from_file(str(PROJECT_ROOT / "app.py")).run()
        app.file_uploader[0].upload(
            "no-table.pdf", make_text_only_pdf(), "application/pdf"
        ).run()
        next(button for button in app.button if button.label == "Start Processing").click().run()

        self.assertEqual(len(app.exception), 0)
        self.assertEqual(len(app.get("download_button")), 0)
        self.assertTrue(
            any("No tables were detected" in info.value for info in app.info)
        )

    def test_scanned_upload_flows_through_ocr_status_and_results(self) -> None:
        fake_ocr_result = {
            "pages": [
                {
                    "page_number": 1,
                    "text": "Mocked OCR scan text",
                    "words": [],
                    "word_count": 4,
                    "status": "ocr_text",
                    "error": None,
                    "engine": "Tesseract",
                    "language": "eng",
                    "dpi": 240,
                    "page_width": 612.0,
                    "page_height": 792.0,
                }
            ],
            "warnings": [],
            "fatal_error": None,
        }
        app = AppTest.from_file(str(PROJECT_ROOT / "app.py")).run()
        app.file_uploader[0].upload("scanned.pdf", make_scanned_pdf(), "application/pdf").run()
        with patch("modules.pdf_extractor.extract_text_with_ocr", return_value=fake_ocr_result):
            next(button for button in app.button if button.label == "Start Processing").click().run()

        self.assertEqual(len(app.exception), 0)
        ocr_metric = next(metric for metric in app.metric if metric.label == "OCR pages · succeeded / needs review")
        self.assertEqual(ocr_metric.value, "1 / 0")
        self.assertTrue(any("no extractable tables" in warning.value.lower() for warning in app.warning))

    def test_failed_ocr_page_shows_runtime_recovery_guidance(self) -> None:
        failed_ocr = {
            "pages": [
                {
                    "page_number": 1,
                    "text": "",
                    "words": [],
                    "word_count": 0,
                    "status": "ocr_failed",
                    "error": "Tesseract was not found. Install it locally.",
                    "engine": "Tesseract",
                    "language": "eng",
                    "dpi": 240,
                    "page_width": 612.0,
                    "page_height": 792.0,
                }
            ],
            "warnings": ["Page 1 OCR failed."],
            "fatal_error": None,
        }
        app = AppTest.from_file(str(PROJECT_ROOT / "app.py")).run()
        app.file_uploader[0].upload("retry.pdf", make_scanned_pdf(), "application/pdf").run()
        with patch("modules.pdf_extractor.extract_text_with_ocr", return_value=failed_ocr):
            next(button for button in app.button if button.label == "Start Processing").click().run()

        self.assertEqual(len(app.exception), 0)
        retry_button = next(button for button in app.button if button.label == "Retry selected OCR pages")
        self.assertTrue(retry_button.disabled)
        self.assertTrue(any("OCR retry is unavailable" in info.value for info in app.info))
        self.assertEqual(app.session_state["extraction_result"]["text"]["pages"][0]["ocr_status"], "ocr_failed")

    def test_retry_failure_remains_visible_and_keeps_prior_success(self) -> None:
        def ocr_page(page_number: int, text: str, status: str, error: str | None = None) -> dict[str, object]:
            return {
                "page_number": page_number,
                "text": text,
                "words": [],
                "word_count": len(text.split()),
                "status": status,
                "error": error,
                "engine": "Tesseract",
                "language": "eng",
                "dpi": 240,
                "page_width": 612.0,
                "page_height": 792.0,
            }

        fake_ocr = {
            "pages": [ocr_page(1, "Good page", "ocr_text"), ocr_page(2, "", "ocr_failed", "Initial failure")],
            "warnings": ["Page 2: initial OCR failure."],
            "fatal_error": None,
        }
        app = AppTest.from_file(str(PROJECT_ROOT / "app.py")).run()
        app.file_uploader[0].upload("retry-failure.pdf", make_scanned_pdf(2), "application/pdf").run()
        with patch(
            "modules.pdf_extractor.extract_text_with_ocr",
            return_value=fake_ocr,
        ):
            next(button for button in app.button if button.label == "Start Processing").click().run()

        self.assertEqual(len(app.exception), 0)
        retry_button = next(button for button in app.button if button.label == "Retry selected OCR pages")
        self.assertTrue(retry_button.disabled)
        self.assertTrue(any("OCR retry is unavailable" in info.value for info in app.info))
        pages = app.session_state["extraction_result"]["text"]["pages"]
        self.assertEqual(pages[0]["ocr_status"], "ocr_text")
        self.assertEqual(pages[1]["ocr_status"], "ocr_failed")
        self.assertEqual(pages[1]["recovery_status"], "failed")

    def test_export_failure_is_actionable_and_does_not_offer_download(self) -> None:
        app = AppTest.from_file(str(PROJECT_ROOT / "app.py")).run()
        app.file_uploader[0].upload("export-failure.pdf", make_table_pdf(), "application/pdf").run()
        with patch("pandas.ExcelWriter", side_effect=Exception("simulated writer failure")):
            next(button for button in app.button if button.label == "Start Processing").click().run()

        self.assertEqual(len(app.exception), 0)
        self.assertIsNone(app.session_state["extraction_result"]["excel_bytes"])
        self.assertTrue(any("Workbook export failed" in error.value for error in app.error))
        self.assertTrue(any("Recovery:" in info.value for info in app.info))
        self.assertEqual(len(app.get("download_button")), 0)

    def test_validation_report_shows_findings_and_severity_filter(self) -> None:
        app = AppTest.from_file(str(PROJECT_ROOT / "app.py")).run()
        app.file_uploader[0].upload(
            "validation.pdf", make_validation_pdf(), "application/pdf"
        ).run()
        next(button for button in app.button if button.label == "Start Processing").click().run()

        self.assertEqual(len(app.exception), 0)
        result = app.session_state["extraction_result"]["validation"]
        self.assertGreater(len(result["findings"]), 0)
        severity_filter = next(selectbox for selectbox in app.selectbox if "Filter findings" in selectbox.label)
        self.assertEqual(severity_filter.value, "All")
        self.assertTrue(any("numeric_text" in finding["category"] for finding in result["findings"]))
        severity_filter.select("Warning").run()
        self.assertEqual(len(app.exception), 0)


if __name__ == "__main__":
    unittest.main()
