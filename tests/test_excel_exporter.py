"""Tests for in-memory Excel export of extracted PDF tables."""

from __future__ import annotations

from io import BytesIO
import unittest

from openpyxl import load_workbook

from modules.excel_exporter import ExcelExportError, export_tables_to_excel


class ExcelExporterTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tables = [
            {
                "table_id": "page-1/table:1",
                "source_page": 1,
                "column_count": 2,
                "headers": ["Item", "Amount"],
                "headers_inferred": True,
                "rows": [
                    {
                        "cells": ["Item", "Amount"],
                        "source_page": 1,
                        "is_inferred_header": True,
                    },
                    {"cells": ["Invoice A", "125.50"], "source_page": 1},
                    {"cells": ["Partial row", ""], "source_page": 2},
                    {"cells": ["=1+1", "kept literally"], "source_page": 2},
                ],
            },
            {
                "table_id": "page-1/table:1",
                "source_page": 3,
                "column_count": 2,
                "headers": ["Code", "Description"],
                "headers_inferred": True,
                "rows": [
                    {
                        "cells": ["Code", "Description"],
                        "source_page": 3,
                        "is_inferred_header": True,
                    },
                    {"cells": ["X-1", "Second table"], "source_page": 3},
                ],
            },
        ]

    def test_workbook_has_sheets_headers_values_and_page_metadata(self) -> None:
        workbook_bytes = export_tables_to_excel(self.tables)
        workbook = load_workbook(BytesIO(workbook_bytes), data_only=False)

        self.assertEqual(
            workbook.sheetnames,
            ["page-1_table_1", "page-1_table_1_2", "Validation Report", "Source Evidence"],
        )
        first = workbook["page-1_table_1"]
        self.assertEqual(
            [cell.value for cell in first[1]],
            ["Item", "Amount", "Source Page"],
        )
        self.assertEqual(
            [cell.value for cell in first[2]],
            ["Invoice A", "125.50", 1],
        )
        self.assertEqual(
            [cell.value for cell in first[3]],
            ["Partial row", None, 2],
        )
        self.assertEqual(first["A4"].value, "=1+1")
        self.assertEqual(first["A4"].data_type, "s")
        self.assertEqual(
            [cell.value for cell in workbook["page-1_table_1_2"][1]],
            ["Code", "Description", "Source Page"],
        )
        self.assertEqual(workbook["page-1_table_1_2"][2][2].value, 3)

    def test_source_evidence_sheet_exports_cell_traceability_and_unavailable_status(self) -> None:
        table = {
            "table_id": "evidence-table",
            "source_page": 3,
            "column_count": 2,
            "headers": ["ID", "Amount"],
            "headers_inferred": True,
            "rows": [
                {
                    "row_index": 0,
                    "source_row_index": 1,
                    "cells": ["ID", "Amount"],
                    "source_page": 3,
                    "is_inferred_header": True,
                    "extraction_method": "pdfplumber",
                    "cell_evidence": [
                        {
                            "source_page": 3,
                            "source_row": 1,
                            "source_column": 1,
                            "extraction_method": "pdfplumber",
                            "coordinate_system": "PDF points, top-left origin",
                            "bbox": [10, 20, 50, 40],
                            "source_text": "ID",
                            "word_boxes": [{"text": "ID", "bbox": [12, 24, 26, 34]}],
                            "evidence_status": "source_text_mapped",
                            "confidence_or_limitations": "Mapped to cell.",
                        },
                        {
                            "source_page": 3,
                            "source_row": 1,
                            "source_column": 2,
                            "extraction_method": "pdfplumber",
                            "evidence_status": "unavailable",
                            "confidence_or_limitations": "No source text mapping available.",
                        },
                    ],
                },
                {
                    "row_index": 1,
                    "source_row_index": 2,
                    "cells": ["A-7", "12.50"],
                    "source_page": 3,
                    "is_inferred_header": False,
                    "extraction_method": "pdfplumber",
                    "cell_evidence": [
                        {"evidence_status": "unavailable", "source_page": 3},
                        {"evidence_status": "unavailable", "source_page": 3},
                    ],
                },
            ],
        }
        workbook = load_workbook(
            BytesIO(export_tables_to_excel([table])),
            data_only=True,
        )
        worksheet = workbook["Source Evidence"]

        self.assertEqual(worksheet["A1"].value, "Table")
        self.assertEqual(worksheet["A2"].value, "evidence-table")
        self.assertEqual(worksheet["D2"].value, "ID (column 1)")
        self.assertEqual(worksheet["E2"].value, 1)
        self.assertEqual(worksheet["F2"].value, "ID")
        self.assertEqual(worksheet["G2"].value, 3)
        self.assertEqual(worksheet["H2"].value, "pdfplumber")
        self.assertEqual(worksheet["I2"].value, "digital_text")
        self.assertEqual(worksheet["K2"].value, "source_text_mapped")
        self.assertIn("10", worksheet["L2"].value)
        self.assertEqual(worksheet["M2"].value, "ID")
        self.assertEqual(worksheet["K3"].value, "unavailable")
        self.assertEqual(worksheet["O3"].value, "No source text mapping available.")
        self.assertEqual(worksheet["F5"].value, "12.50")
        self.assertEqual(worksheet["G5"].value, 3)

    def test_validation_report_sheet_contains_source_locations(self) -> None:
        report = {
            "summary": "1 validation finding(s) across 2 table(s).",
            "findings": [
                {
                    "table_name": "page-1-table-1",
                    "source_page": 4,
                    "row": 3,
                    "related_row": None,
                    "column": "Amount (column 2)",
                    "severity": "warning",
                    "category": "numeric_text",
                    "issue": "Unexpected numeric text",
                    "suggested_action": "Check the source page.",
                    "evidence": "Value: N/A",
                    "actual_value": "110.00",
                    "expected_value": "105.00",
                }
            ],
        }
        workbook = load_workbook(
            BytesIO(export_tables_to_excel(self.tables, validation_report=report)),
            data_only=True,
        )
        worksheet = workbook["Validation Report"]

        self.assertEqual(worksheet["A1"].value, "Table")
        self.assertEqual(worksheet["B2"].value, 4)
        self.assertEqual(worksheet["C2"].value, 3)
        self.assertEqual(worksheet["E2"].value, "Amount (column 2)")
        self.assertEqual(worksheet["G2"].value, "numeric_text")
        self.assertEqual(worksheet["H2"].value, "Unexpected numeric text")
        self.assertEqual(worksheet["K1"].value, "Actual Value")
        self.assertEqual(worksheet["K2"].value, "110.00")
        self.assertEqual(worksheet["L2"].value, "105.00")

    def test_identifier_date_like_and_signed_decimal_values_are_preserved(self) -> None:
        table = {
            "table_id": "precision-check",
            "source_page": 6,
            "column_count": 4,
            "headers": ["Identifier", "Date", "Amount", "Discount"],
            "headers_inferred": True,
            "rows": [
                {
                    "cells": ["Identifier", "Date", "Amount", "Discount"],
                    "source_page": 6,
                    "is_inferred_header": True,
                },
                {"cells": ["0006", "01/02/2026", "-125.375", "0.005"], "source_page": 6},
            ],
        }

        workbook = load_workbook(
            BytesIO(export_tables_to_excel([table])),
            data_only=True,
        )
        worksheet = workbook["precision-check"]

        self.assertEqual(worksheet.max_row, 2)
        self.assertEqual(worksheet.max_column, 5)
        self.assertEqual(worksheet["A2"].value, "0006")
        self.assertEqual(worksheet["B2"].value, "01/02/2026")
        self.assertEqual(worksheet["C2"].value, "-125.375")
        self.assertEqual(worksheet["D2"].value, "0.005")
        self.assertEqual(worksheet["E2"].value, 6)

    def test_no_issue_workbook_contains_explicit_report_message_and_keeps_original_values(self) -> None:
        original_amount = self.tables[0]["rows"][1]["cells"][1]
        report = {
            "findings": [],
            "summary": "No validation issues were detected by the implemented checks.",
        }

        workbook = load_workbook(
            BytesIO(export_tables_to_excel(self.tables, validation_report=report)),
            data_only=True,
        )

        self.assertIn("No validation issues were detected", workbook["Validation Report"]["A2"].value)
        self.assertEqual(workbook["page-1_table_1"]["B2"].value, original_amount)
        self.assertEqual(workbook["page-1_table_1"]["A3"].value, "Partial row")

    def test_empty_table_list_is_a_helpful_export_error(self) -> None:
        with self.assertRaisesRegex(ExcelExportError, "No tables were detected"):
            export_tables_to_excel([])

    def test_table_with_no_rows_or_columns_is_rejected(self) -> None:
        with self.assertRaisesRegex(ExcelExportError, "is empty"):
            export_tables_to_excel([{"table_id": "empty", "rows": []}])


if __name__ == "__main__":
    unittest.main()
