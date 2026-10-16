"""Build reproducible challenge PDFs and process them through the real pipeline.

Run from the repository root with the project's virtual environment active:
    python samples/build_demo_samples.py

PDF fixtures are generated with known, local content. Excel workbooks are emitted
only when the actual extraction pipeline detects at least one table. In particular,
the scanned fixture does not receive fabricated OCR output when Tesseract is absent.
"""

from __future__ import annotations

import json
from pathlib import Path
import sys
from typing import Any

from openpyxl import load_workbook
from io import BytesIO

try:
    import pymupdf as fitz
except ImportError:  # Compatibility with earlier PyMuPDF versions.
    import fitz
from PIL import Image, ImageDraw, ImageFont

PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = PROJECT_ROOT / "samples" / "outputs"
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from modules.excel_exporter import export_tables_to_excel
from modules.pdf_extractor import extract_pdf_tables, extract_pdf_text
from modules.table_stitcher import stitch_tables
from modules.validator import validate_tables

SAMPLES_DIR = PROJECT_ROOT / "samples"


def _draw_ruled_table(
    page: fitz.Page,
    rows: list[list[str]],
    x_edges: list[float],
    y_top: float,
    row_height: float,
    font_size: float = 8,
) -> None:
    """Draw table rules and selectable text into a PDF page."""
    y_bottom = y_top + row_height * len(rows)
    for x in x_edges:
        page.draw_line((x, y_top), (x, y_bottom), color=(0, 0, 0), width=0.65)
    for row_index in range(len(rows) + 1):
        y = y_top + row_index * row_height
        page.draw_line((x_edges[0], y), (x_edges[-1], y), color=(0, 0, 0), width=0.65)
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            if value:
                page.insert_text(
                    (x_edges[column_index] + 3, y_top + row_index * row_height + row_height * 0.68),
                    value,
                    fontsize=font_size,
                )


def _write_ten_column_digital_pdf(path: Path) -> list[list[str]]:
    """Create one digital page with a ruled, ten-column known table."""
    headers = [f"Field {index:02}" for index in range(1, 11)]
    first_data = ["ID-001", "-12.50", "1,234.56", "2026-10-16", "North", "06", "7.00", "8.00", "9.00", "10.00"]
    second_data = ["ID-002", "0.00", "250.00", "2026-10-17", "South", "11", "12.00", "13.00", "14.00", "15.00"]
    rows = [headers, first_data, second_data]
    document = fitz.open()
    page = document.new_page(width=792, height=612)
    x_edges = [24 + index * 74 for index in range(11)]
    _draw_ruled_table(page, rows, x_edges, 150, 48, font_size=7)
    document.save(path)
    document.close()
    return rows


def _write_scanned_pdf(path: Path) -> list[list[str]]:
    """Create an image-only table page; OCR output is not simulated here."""
    rows = [
        ["Reference", "Date", "Amount"],
        ["SCAN-101", "2026-10-16", "-25.75"],
        ["SCAN-102", "2026-10-17", "310.00"],
    ]
    image = Image.new("RGB", (1275, 1650), "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("arial.ttf", 32)
    except OSError:
        font = ImageFont.load_default()
    x_edges = [100, 480, 850, 1170]
    y_top, row_height = 270, 110
    for x in x_edges:
        draw.line((x, y_top, x, y_top + row_height * len(rows)), fill="black", width=3)
    for index in range(len(rows) + 1):
        y = y_top + index * row_height
        draw.line((x_edges[0], y, x_edges[-1], y), fill="black", width=3)
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            draw.text(
                (x_edges[column_index] + 12, y_top + row_index * row_height + 32),
                value,
                fill="black",
                font=font,
            )
    image_path = path.with_suffix(".png")
    image.save(image_path)
    document = fitz.open()
    page = document.new_page(width=612, height=792)
    page.insert_image(page.rect, filename=str(image_path))
    document.save(path)
    document.close()
    image_path.unlink()
    return rows


def _write_multipage_pdf(path: Path) -> list[list[list[str]]]:
    """Create a two-page digital table with a repeated page-two header."""
    headers = ["Account Code", "Description", "Amount"]
    page_one_rows = [headers, ["AC-01", "Opening balance", "100.00"], ["AC-02", "Service fee", "-5.00"]]
    page_two_rows = [headers, ["AC-03", "Customer payment", "75.00"], ["AC-04", "Closing balance", "170.00"]]
    document = fitz.open()
    for rows, y_top in ((page_one_rows, 600), (page_two_rows, 42)):
        page = document.new_page(width=612, height=792)
        _draw_ruled_table(page, rows, [70, 230, 430, 542], y_top, 36, font_size=8)
    document.save(path)
    document.close()
    return [page_one_rows, page_two_rows]


def _run_real_pipeline(pdf_path: Path, expected_rows: list[list[list[str]]] | list[list[str]]) -> dict[str, Any]:
    """Run actual text/OCR/table/stitch/validation/export functions for a sample."""
    pdf_bytes = pdf_path.read_bytes()
    text_result = extract_pdf_text(pdf_bytes, filename=pdf_path.name)
    table_result = extract_pdf_tables(
        pdf_bytes,
        filename=pdf_path.name,
        ocr_pages=text_result["pages"],
    )
    stitch_result = stitch_tables(table_result["tables"])
    validation_result = validate_tables(stitch_result.tables)

    tables = stitch_result.tables
    if pdf_path.stem == "digital_10_column":
        assert tables, "Digital 10-column fixture produced no detected tables."
        assert any(table["column_count"] >= 10 for table in tables), "Detected table has fewer than 10 columns."
        cell_values = [cell for table in tables for row in table["rows"] for cell in row["cells"]]
        assert "ID-001" in cell_values and "-12.50" in cell_values and "1,234.56" in cell_values
    elif pdf_path.stem == "multipage_continuation":
        assert tables, "Multi-page fixture produced no detected tables."
        assert stitch_result.stitched_count >= 1, "Expected page continuation was not stitched."
        merged = max(tables, key=lambda table: int(table.get("row_count", 0)))
        all_rows = [row["cells"] for row in merged["rows"]]
        assert all_rows.count(["Account Code", "Description", "Amount"]) == 1, "Repeated header was not removed."
        assert any(row.get("source_page") == 1 for row in merged["rows"])
        assert any(row.get("source_page") == 2 for row in merged["rows"])

    workbook_path = OUTPUT_DIR / f"{pdf_path.stem}.xlsx" if tables else None
    workbook_bytes = None
    export_error = None
    workbook_sheets: list[str] = []
    if tables:
        workbook_bytes = export_tables_to_excel(tables, validation_report=validation_result)
        workbook_path.write_bytes(workbook_bytes)
        workbook = load_workbook(BytesIO(workbook_bytes), data_only=True)
        workbook_sheets = list(workbook.sheetnames)
        assert "Validation Report" in workbook_sheets
        assert "Source Evidence" in workbook_sheets
        data_worksheets = [
            workbook[name]
            for name in workbook_sheets
            if name not in {"Validation Report", "Source Evidence"}
        ]
        assert len(data_worksheets) == len(tables)
        if pdf_path.stem == "digital_10_column":
            worksheet = data_worksheets[0]
            assert worksheet.max_column == 11  # Ten data columns and source page.
            assert worksheet.cell(row=1, column=1).value == "Field 01"
            assert worksheet.cell(row=2, column=1).value == "ID-001"
            assert worksheet.cell(row=2, column=2).value == "-12.50"
            assert worksheet.cell(row=2, column=11).value == 1
        elif pdf_path.stem == "multipage_continuation":
            worksheet = data_worksheets[0]
            data_row_count = worksheet.max_row - 1
            identifiers = [worksheet.cell(row=index, column=1).value for index in range(2, worksheet.max_row + 1)]
            source_pages = [worksheet.cell(row=index, column=4).value for index in range(2, worksheet.max_row + 1)]
            assert data_row_count == 4, "Workbook must contain four data rows and a single header row."
            assert identifiers == ["AC-01", "AC-02", "AC-03", "AC-04"]
            assert source_pages == [1, 1, 2, 2]

    return {
        "pdf": str(pdf_path.relative_to(PROJECT_ROOT)),
        "page_status": [
            {
                "page": page["page_number"],
                "digital_text": page["has_digital_text"],
                "ocr_status": page["ocr_status"],
                "error": page["error"],
            }
            for page in text_result["pages"]
        ],
        "detected_table_count": len(table_result["tables"]),
        "stitched_table_count": stitch_result.stitched_count,
        "tables": [
            {
                "name": table["table_id"],
                "columns": table["column_count"],
                "rows": table["row_count"],
                "source_pages": table.get("source_pages", [table.get("source_page")]),
                "headers": table.get("headers"),
                "extraction_method": table.get("extraction_method", "pdfplumber"),
                "data_rows": [row["cells"] for row in table["rows"]],
                "source_page_per_row": [row.get("source_page") for row in table["rows"]],
            }
            for table in tables
        ],
        "validation_summary": validation_result["summary"],
        "validation_findings": validation_result["findings"],
        "source_evidence_available_cells": sum(
            1
            for table in tables
            for row in table["rows"]
            for evidence in row.get("cell_evidence", [])
            if evidence.get("evidence_status") in {"source_text_mapped", "cell_region_only"}
        ),
        "workbook": str(workbook_path.relative_to(PROJECT_ROOT)) if workbook_path else None,
        "workbook_sheets": workbook_sheets,
        "export_error": export_error,
        "limitations": (
            "Real OCR requires the separately installed Tesseract application and eng trained data."
            if any(page["ocr_status"] in {"ocr_failed", "ocr_no_text"} for page in text_result["pages"])
            else None
        ),
    }


def main() -> None:
    """Create all fixtures, run each through real pipeline code, and save a manifest."""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    digital_pdf = SAMPLES_DIR / "digital_10_column.pdf"
    scanned_pdf = SAMPLES_DIR / "scanned_table.pdf"
    multipage_pdf = SAMPLES_DIR / "multipage_continuation.pdf"

    digital_rows = _write_ten_column_digital_pdf(digital_pdf)
    scanned_rows = _write_scanned_pdf(scanned_pdf)
    multipage_rows = _write_multipage_pdf(multipage_pdf)

    records = [
        _run_real_pipeline(digital_pdf, digital_rows),
        _run_real_pipeline(scanned_pdf, scanned_rows),
        _run_real_pipeline(multipage_pdf, multipage_rows),
    ]
    manifest = {
        "generator": "samples/build_demo_samples.py",
        "claim": "All listed extraction and workbook results are generated by the project pipeline at run time; no OCR values are mocked.",
        "samples": records,
    }
    manifest_path = OUTPUT_DIR / "sample_run_manifest.json"
    manifest_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(manifest, indent=2, ensure_ascii=False))
    print(f"Manifest written to {manifest_path.relative_to(PROJECT_ROOT)}")


if __name__ == "__main__":
    main()
